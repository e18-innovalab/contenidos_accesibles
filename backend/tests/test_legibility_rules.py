"""Reglas de legibilidad visual: letra chica (issue #79, parte 1). Los archivos se generan en el test."""

import io

import pymupdf
import pytest
from docx import Document
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches
from pptx.util import Pt as PptPt

from app.schemas.document import ExtractedDocument, LegibilityStats, SmallTextPassage
from app.services import rules_service as rs
from app.services.document_service import document_service
from app.services.legibility_rules import SMALL_BODY_POINTS
from app.services.office_service import office_service

BODY = (
    "Los ecosistemas acuáticos reúnen seres vivos que dependen del agua para vivir. "
    "En esta actividad vamos a observar cómo se relacionan entre sí y con su ambiente. "
)
NOTE = (
    "Nota para el docente: esta actividad puede extenderse a dos clases si el grupo necesita más tiempo "
    "para observar las muestras, registrar los datos en la tabla y comparar los resultados con el resto."
)


def _legibility(result: rs.RulesResult) -> list[dict]:
    return [f.barrier for f in result.findings if "letra" in f.barrier["issue"].lower()]


def _penalty_rules(result: rs.RulesResult) -> set[str]:
    return {p.rule for p in result.penalties}


def _visual(result: rs.RulesResult):
    return next(d for d in result.dimensions if d.category == "accesibilidad_visual")


# --- Generadores de archivos -----------------------------------------------


def _pdf(*boxes: tuple[str, float]) -> bytes:
    """Un cuadro de texto por (texto, tamaño), uno debajo del otro."""
    doc = pymupdf.open()
    page = doc.new_page()
    y = 50
    for text, size in boxes:
        height = 400 if size >= 9 else 120
        page.insert_textbox(pymupdf.Rect(50, y, 550, y + height), text, fontsize=size)
        y += height + 10
    data = doc.tobytes()
    doc.close()
    return data


def _docx(*paragraphs: tuple[str, float | None], normal_size: float | None = None) -> bytes:
    document = Document()
    if normal_size is not None:
        document.styles["Normal"].font.size = Pt(normal_size)
    for text, size in paragraphs:
        run = document.add_paragraph().add_run(text)
        if size is not None:
            run.font.size = Pt(size)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pptx(*paragraphs: tuple[str, float | None], font_scale: str | None = None) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # en blanco
    frame = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(6)).text_frame
    for index, (text, size) in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = text
        if size is not None:
            run.font.size = PptPt(size)
    if font_scale:
        body_pr = frame._bodyPr
        autofit = body_pr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}normAutofit", {})
        autofit.set("fontScale", font_scale)
        body_pr.append(autofit)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _rules_pdf(*boxes) -> rs.RulesResult:
    return rs.run_rules(document_service.process_pdf(_pdf(*boxes), "material.pdf"))


def _rules_docx(*paragraphs, **kwargs) -> rs.RulesResult:
    return rs.run_rules(office_service.process_docx(_docx(*paragraphs, **kwargs), "material.docx"))


def _rules_pptx(*paragraphs, **kwargs) -> rs.RulesResult:
    return rs.run_rules(office_service.process_pptx(_pptx(*paragraphs, **kwargs), "material.pptx"))


# --- Cuerpo del texto en letra chica -----------------------------------------


@pytest.mark.parametrize("run_rules, size", [
    (lambda s: _rules_pdf((BODY * 6, s)), 6),
    (lambda s: _rules_docx((BODY * 6, s)), 6),
    (lambda s: _rules_pptx((BODY * 3, s)), 6),
    (lambda s: _rules_pdf((BODY * 6, s)), 8.5),
])
def test_cuerpo_en_letra_chica_es_barrera(run_rules, size):
    result = run_rules(size)

    barriers = _legibility(result)
    assert len(barriers) == 1
    assert barriers[0]["severity"] == "barrera_detectada"
    assert barriers[0]["category"] == "accesibilidad_visual"
    assert barriers[0]["source"] == "rules"
    assert "11 pt" in barriers[0]["recommendation"]
    assert _visual(result).status == "barrera_detectada"
    assert "letra_chica_cuerpo" in _penalty_rules(result)
    assert result.metrics.body_font_size_pt == size
    assert rs.compute_score(None, result).rules_score < 100


def test_criterio_de_aceptacion_cuerpo_6pt_ya_no_saca_100():
    # El mismo material en 6 pt pierde la penalización de letra chica respecto de 11 pt y queda con tope.
    clean = rs.compute_score(None, _rules_pdf((BODY * 6, 11)))
    small = rs.compute_score(None, _rules_pdf((BODY * 6, 6)))
    assert small.rules_score == clean.rules_score - SMALL_BODY_POINTS
    assert small.final_score <= rs.SEVERE_BARRIER_CAP


def test_word_detecta_letra_chica_heredada_del_estilo_normal():
    # Ningún run define tamaño: lo hereda del estilo Normal.
    result = _rules_docx((BODY * 6, None), normal_size=7)
    assert [b["severity"] for b in _legibility(result)] == ["barrera_detectada"]


def test_powerpoint_aplica_la_reduccion_automatica_del_texto():
    # 14 pt reducidos al 50% por "Reducir el texto al desbordar" se ven en 7 pt.
    result = _rules_pptx((BODY * 3, 14), font_scale="50000")
    assert result.metrics.body_font_size_pt == 7
    assert [b["severity"] for b in _legibility(result)] == ["barrera_detectada"]


# --- Pasajes largos en letra chica (notas, epígrafes) -------------------------


def test_pasaje_largo_en_letra_chica_en_pdf():
    result = _rules_pdf((BODY * 6, 11), (NOTE, 7))

    barriers = _legibility(result)
    assert len(barriers) == 1
    assert barriers[0]["severity"] == "para_revisar"
    assert "(página 1)" in barriers[0]["issue"]
    assert barriers[0]["original_text"].startswith("Nota para el docente")
    assert barriers[0]["fragment_found"] is True
    assert _visual(result).status == "para_revisar"
    assert "letra_chica_pasajes" in _penalty_rules(result)
    assert result.metrics.small_text_passages == 1


def test_pasaje_largo_en_letra_chica_en_word_y_powerpoint():
    docx = _legibility(_rules_docx((BODY * 6, 11), (NOTE, 7)))
    pptx = _legibility(_rules_pptx((BODY * 3, 20), (NOTE, 7)))

    assert [b["severity"] for b in docx] == ["para_revisar"]
    assert "página" not in docx[0]["issue"]  # Word no informa páginas
    assert [b["severity"] for b in pptx] == ["para_revisar"]
    assert "(diapositiva 1)" in pptx[0]["issue"]


def test_se_listan_hasta_tres_pasajes_y_la_penalizacion_tiene_tope():
    # El cuerpo tiene que seguir siendo mayoría: si no, el tamaño predominante pasaría a ser 7 pt.
    result = _rules_docx(*[(BODY * 6, 11)] * 4, *[(NOTE, 7)] * 6)

    assert len(_legibility(result)) == 3
    penalty = next(p for p in result.penalties if p.rule == "letra_chica_pasajes")
    assert penalty.count == 6
    assert penalty.points == 15


def test_si_todo_el_cuerpo_es_chico_no_se_repite_por_pasaje():
    result = _rules_docx(*[(NOTE, 6)] * 4)
    assert [b["severity"] for b in _legibility(result)] == ["barrera_detectada"]
    assert "letra_chica_pasajes" not in _penalty_rules(result)


# --- Casos que no deben marcar nada -------------------------------------------


@pytest.mark.parametrize("run_rules", [
    lambda: _rules_pdf((BODY * 6, 11)),
    lambda: _rules_pdf((BODY * 6, 12)),
    lambda: _rules_docx((BODY * 6, None)),  # tamaño por defecto de Word (11 pt)
    lambda: _rules_docx((BODY * 6, 12)),
    lambda: _rules_pptx((BODY * 3, 18)),
])
def test_letra_de_11_12_pt_no_genera_falsos_positivos(run_rules):
    result = run_rules()
    assert _legibility(result) == []
    assert not {"letra_chica_cuerpo", "letra_chica_pasajes"} & _penalty_rules(result)


def test_texto_corto_en_letra_chica_no_es_pasaje():
    # Número de página, pie de foto o fuente: menos de 20 palabras.
    result = _rules_pdf((BODY * 6, 11), ("Página 1 de 3", 6), ("Fuente: elaboración propia", 7))
    assert _legibility(result) == []


def test_powerpoint_sin_tamanos_explicitos_no_marca_nada():
    # El tamaño se hereda del patrón de diapositivas: no se conoce, no se inventa.
    extracted = office_service.process_pptx(_pptx((BODY * 3, None)), "material.pptx")
    assert extracted.legibility.measured_chars == 0
    assert _legibility(rs.run_rules(extracted)) == []


def test_texto_pegado_no_se_evalua():
    text = BODY * 6
    doc = ExtractedDocument(
        source_type="text",
        raw_text=text,
        text_for_analysis=text,
        word_count=rs.count_words(text),
        character_count=len(text),
        # Aunque llegaran datos tipográficos, el texto pegado no tiene tipografía.
        legibility=LegibilityStats(
            body_font_size=6, measured_chars=len(text),
            small_text_passages=[SmallTextPassage(text=NOTE, font_size=6, word_count=30)],
            small_text_passages_total=1,
        ),
    )
    result = rs.run_rules(doc)
    assert _legibility(result) == []
    assert result.metrics.body_font_size_pt is None


def test_texto_pegado_por_document_service_no_trae_datos_tipograficos():
    assert document_service.process_text(BODY * 3).legibility is None


def test_pdf_escaneado_no_se_evalua_con_letra_chica():
    doc = document_service.process_pdf(_pdf((BODY * 6, 6)), "material.pdf").model_copy(
        update={"is_scanned_pdf": True}
    )
    assert _legibility(rs.run_rules(doc)) == []
