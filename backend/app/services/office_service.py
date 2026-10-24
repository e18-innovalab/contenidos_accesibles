"""
Extracción de Word (.docx) y PowerPoint (.pptx) al mismo ExtractedDocument que un PDF (Issue #47).

A diferencia del PDF, estos formatos guardan la estructura de forma explícita: los títulos
salen de los estilos (Título 1, Título 2...) o del título de cada diapositiva, las listas
de la numeración/viñetas del párrafo y las imágenes traen su texto alternativo (o la marca
de decorativa). Las imágenes que ya lo tienen no se envían a la IA para proponer uno.
Todo se procesa en memoria, igual que el PDF.
"""
import io
import re
import zipfile
from dataclasses import dataclass, field

import pymupdf
from docx import Document as open_docx
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn as w_qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image as PILImage
from pptx import Presentation as open_pptx
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn as p_qn
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture
from pptx.util import Emu

from app.schemas.document import DocumentStructure, ExtractedDocument, ExtractedImage, HeadingItem
from app.services.document_service import (
    MAX_ASPECT_RATIO,
    MAX_IMAGES_FOR_ALT_TEXT,
    MIN_ASPECT_RATIO,
    MIN_IMAGE_DIMENSION_PX,
    MIN_IMAGE_RENDER_PT,
    DocumentService,
)
from app.services.legibility_extractor import extract_docx_legibility, extract_pptx_legibility

ZIP_MAGIC = b"PK\x03\x04"
# Compound File Binary: .doc/.ppt antiguos y también .docx/.pptx protegidos con contraseña.
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

# Defensa contra zip bombs: el archivo subido pesa como máximo 10 MB, pero descomprimido podría
# ocupar mucho más. Se valida lo declarado en el índice del ZIP antes de abrirlo.
MAX_UNCOMPRESSED_BYTES = 150 * 1024 * 1024
MAX_ZIP_ENTRIES = 5000
# Imágenes de más de 40 megapíxeles no se decodifican (unos 160 MB en memoria).
MAX_IMAGE_PIXELS = 40_000_000

_MAIN_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml": "docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml": "pptx",
}
_MACRO_CONTENT_TYPES = {
    "application/vnd.ms-word.document.macroEnabled.main+xml": ".docm",
    "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml": ".pptm",
}

_FORMAT_LABELS = {"docx": "Word", "pptx": "PowerPoint"}

# Marca de "decorativa" de Office (Formato de imagen → Texto alternativo → Marcar como decorativa).
_DECORATIVE_TAG = "{http://schemas.microsoft.com/office/drawing/2017/decorative}decorative"

_HEADING_STYLE_RE = re.compile(r"^(?:heading|t[íi]tulo)\s*(\d)$", re.IGNORECASE)
_FILENAME_ALT_RE = re.compile(r"^[\w\-. ]+\.(?:png|jpe?g|gif|bmp|tiff?|emf|wmf|svg|webp)$", re.IGNORECASE)
_PAGES_RE =re.compile(rb"<(?:\w+:)?Pages>(\d+)</(?:\w+:)?Pages>")
_EMU_PER_PT = 12700


def detect_office_format(file_bytes: bytes) -> str:
    """
    Devuelve "docx" o "pptx" según el contenido real del archivo (no la extensión).

    Lanza ValueError con un mensaje para el docente si el archivo está protegido, es un
    formato antiguo, tiene macros, no es un documento de Office o descomprimido es demasiado grande.
    """
    if file_bytes.startswith(OLE_MAGIC):
        raise ValueError(
            "El archivo está protegido con contraseña o es un formato antiguo de Office (.doc/.ppt). "
            "Guardalo como .docx o .pptx sin contraseña y volvé a subirlo."
        )
    if not file_bytes.startswith(ZIP_MAGIC):
        raise ValueError("El contenido del archivo no corresponde a un documento de Word o PowerPoint válido.")

    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
            entries = zf.infolist()
            if len(entries) > MAX_ZIP_ENTRIES or sum(e.file_size for e in entries) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("El documento es demasiado grande una vez descomprimido para poder analizarlo.")
            try:
                content_types = zf.read("[Content_Types].xml").decode("utf-8", errors="replace")
            except KeyError:
                raise ValueError(
                    "El contenido del archivo no corresponde a un documento de Word o PowerPoint válido."
                ) from None
    except zipfile.BadZipFile as e:
        raise ValueError("El archivo de Word o PowerPoint está dañado y no se puede abrir.") from e

    for content_type, extension in _MACRO_CONTENT_TYPES.items():
        if content_type in content_types:
            raise ValueError(
                f"Los documentos con macros ({extension}) no se admiten. "
                "Guardalo como .docx o .pptx y volvé a subirlo."
            )
    for content_type, source_type in _MAIN_CONTENT_TYPES.items():
        if content_type in content_types:
            return source_type
    raise ValueError("Formato no admitido. Se aceptan documentos de Word (.docx) y PowerPoint (.pptx).")


# ---------------------------------------------------------------------------
# Piezas comunes a Word y PowerPoint
# ---------------------------------------------------------------------------

@dataclass
class _Block:
    """Unidad de texto en orden de lectura: párrafo, título, ítem de lista o tabla."""
    text: str
    is_list: bool = False
    # Solo PowerPoint: posición vertical (pt) del cuadro de texto, para el contexto de las imágenes.
    top: float = 0.0
    bottom: float = 0.0


@dataclass
class _Images:
    """Resultado de recorrer las imágenes del documento."""
    total: int = 0
    with_alt_text: int = 0
    relevant_without_alt: int = 0
    extracted: list[ExtractedImage] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _join_blocks(blocks: list[_Block]) -> str:
    """Párrafos separados por línea en blanco; los ítems de una misma lista, por un salto simple."""
    parts: list[str] = []
    for i, block in enumerate(blocks):
        if i:
            parts.append("\n" if block.is_list and blocks[i - 1].is_list else "\n\n")
        parts.append(block.text)
    return "".join(parts).strip()


def _context_snippet(preceding: str | None, following: str | None) -> str:
    """Mismo formato que el contexto de las imágenes de un PDF: anterior [...] posterior, hasta 120 palabras."""
    snippet = " [...] ".join(p.strip() for p in (preceding, following) if p and p.strip())
    words = snippet.split()
    if len(words) > 120:
        snippet = " ".join(words[:120]) + "..."
    return snippet


def _alt_text_state(c_nv_pr) -> tuple[bool, bool]:
    """(tiene texto alternativo, está marcada como decorativa) a partir del cNvPr/docPr de la imagen."""
    if c_nv_pr is None:
        return False, False
    descr = (c_nv_pr.get("descr") or "").strip()
    # Algunas herramientas guardan el nombre del archivo ("image1.png") como descripción: no describe nada.
    has_alt = bool(descr) and not _FILENAME_ALT_RE.match(descr)
    decorative = any(el.get("val") in ("1", "true") for el in c_nv_pr.iter(_DECORATIVE_TAG))
    return has_alt, decorative


def _decode_image(blob: bytes) -> tuple[bytes, int, int] | None:
    """PNG para la IA de visión, o None si el formato no se puede decodificar (EMF, WMF, SVG...) o es enorme."""
    try:
        # Pillow lee solo el encabezado: evita decodificar una imagen chica en bytes pero gigante en píxeles.
        with PILImage.open(io.BytesIO(blob)) as header:
            width, height = header.size
        if width * height > MAX_IMAGE_PIXELS:
            return None
        return DocumentService._encode_pixmap_png(pymupdf.Pixmap(blob))
    except Exception:
        return None


def _is_relevant_size(width_pt: float, height_pt: float) -> bool:
    """Mismo filtro que en el PDF: descarta íconos, viñetas gráficas y líneas decorativas."""
    if width_pt < MIN_IMAGE_RENDER_PT or height_pt < MIN_IMAGE_RENDER_PT:
        return False
    ratio = width_pt / max(height_pt, 1)
    return MIN_ASPECT_RATIO <= ratio <= MAX_ASPECT_RATIO


def _add_candidate(
    images: _Images,
    blob: bytes | None,
    page_number: int,
    bbox: tuple[float, float, float, float],
    location_desc: str,
    context: str,
) -> None:
    """Registra una imagen relevante sin texto alternativo y la codifica para la IA (hasta el tope)."""
    images.relevant_without_alt += 1
    if len(images.extracted) >= MAX_IMAGES_FOR_ALT_TEXT:
        return
    encoded = _decode_image(blob) if blob else None
    if encoded is None:
        images.warnings.append(
            f"No se pudo procesar automáticamente la imagen ({location_desc}); "
            "revisá su texto alternativo manualmente."
        )
        return
    img_bytes, width, height = encoded
    if width < MIN_IMAGE_DIMENSION_PX or height < MIN_IMAGE_DIMENSION_PX:
        images.relevant_without_alt -= 1
        return
    images.extracted.append(
        ExtractedImage(
            image_index=len(images.extracted) + 1,
            page_number=page_number,
            bbox=tuple(round(v, 1) for v in bbox),
            width=width,
            height=height,
            location_description=location_desc,
            context_text=context,
            image_bytes=img_bytes,
            mime_type="image/png",
        )
    )


def _build_document(
    source_type: str,
    filename: str,
    blocks: list[_Block],
    page_count: int,
    images: _Images,
    structure: DocumentStructure,
) -> ExtractedDocument:
    full_text = _join_blocks(blocks)
    words = re.findall(r"\b\w+\b", full_text)
    word_count = len(words)
    text_for_analysis, is_truncated, length_warnings = DocumentService._prepare_analysis_text(full_text, words)
    warnings = list(length_warnings) + images.warnings

    if images.relevant_without_alt > MAX_IMAGES_FOR_ALT_TEXT:
        warnings.append(
            f"El documento contiene {images.relevant_without_alt} imágenes relevantes sin texto alternativo; "
            f"se generaron propuestas para las primeras {MAX_IMAGES_FOR_ALT_TEXT} para preservar cuota."
        )
    if word_count == 0:
        warnings.append(
            f"El documento de {_FORMAT_LABELS[source_type]} no contiene texto: "
            "puede que todo el contenido esté dentro de imágenes."
        )
    if images.total:
        pending = images.total - images.with_alt_text
        warnings.append(
            f"Se detectaron {images.total} imagen(es) en el documento; "
            f"{images.with_alt_text} ya tiene(n) texto alternativo o está(n) marcada(s) como decorativa(s)"
            + (f" y {pending} no." if pending else ".")
        )

    return ExtractedDocument(
        source_type=source_type,
        filename=filename,
        raw_text=full_text,
        text_for_analysis=text_for_analysis,
        is_truncated=is_truncated,
        page_count=max(page_count, 1),
        word_count=word_count,
        character_count=len(full_text),
        has_images=images.total > 0,
        image_count=images.total,
        images_with_alt_text=images.with_alt_text,
        is_scanned_pdf=False,
        warnings=warnings,
        structure=structure,
        extracted_images=images.extracted,
    )


def _structure(headings: list[HeadingItem], list_items: int, body_font_size: float | None) -> DocumentStructure:
    return DocumentStructure(
        has_headings=bool(headings),
        body_font_size=body_font_size,
        headings=headings,
        hierarchy_levels=len({h.level for h in headings}),
        list_items_count=list_items,
    )


def _wrap_errors(label: str, func, *args) -> ExtractedDocument:
    try:
        return func(*args)
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"No se pudo procesar el archivo de {label}: {e}") from e


# ---------------------------------------------------------------------------
# Word (.docx)
# ---------------------------------------------------------------------------

def _style_chain(style):
    while style is not None:
        yield style
        style = style.base_style


def _outline_level(p_pr) -> int | None:
    """Nivel de esquema de Word (0 = Título 1). 9 o ausente = texto independiente."""
    if p_pr is None:
        return None
    el = p_pr.find(w_qn("w:outlineLvl"))
    if el is None:
        return None
    try:
        level = int(el.get(w_qn("w:val")))
    except (TypeError, ValueError):
        return None
    return level + 1 if 0 <= level < 9 else None


def _docx_heading_level(paragraph: Paragraph) -> int | None:
    """Nivel de título por el nivel de esquema del párrafo o por su estilo (y los estilos en que se basa)."""
    level = _outline_level(paragraph._p.pPr)
    if level:
        return level
    for style in _style_chain(paragraph.style):
        name = (style.name or "").strip()
        if name.lower() in ("title", "título", "titulo"):
            return 1
        match = _HEADING_STYLE_RE.match(name)
        if match:
            return int(match.group(1))
        level = _outline_level(style.element.pPr)
        if level:
            return level
    return None


def _docx_is_list(paragraph: Paragraph) -> bool:
    """Viñeta o numeración aplicada al párrafo o heredada de su estilo (Lista con viñetas, etc.)."""
    p_pr = paragraph._p.pPr
    if p_pr is not None and p_pr.find(w_qn("w:numPr")) is not None:
        return True
    return any(
        s.element.pPr is not None and s.element.pPr.find(w_qn("w:numPr")) is not None
        for s in _style_chain(paragraph.style)
    )


def _docx_font_size(paragraph: Paragraph) -> float | None:
    for run in paragraph.runs:
        if run.font.size is not None:
            return run.font.size.pt
    for style in _style_chain(paragraph.style):
        if style.font.size is not None:
            return style.font.size.pt
    return None


def _docx_is_bold(paragraph: Paragraph) -> bool:
    """Negrita en los runs con texto o, si no la definen, heredada del estilo."""
    runs = [r for r in paragraph.runs if r.text.strip()]
    if runs and all(r.bold is not None for r in runs):
        return all(r.bold for r in runs)
    return any(s.font.bold for s in _style_chain(paragraph.style))


def _docx_drawings(element) -> list:
    """Imágenes, gráficos y SmartArt del bloque (los cuadros de texto y formas no cuentan)."""
    drawings = []
    for drawing in element.iter(w_qn("w:drawing")):
        graphic_data = next(drawing.iter(w_qn("a:graphicData")), None)
        uri = graphic_data.get("uri", "") if graphic_data is not None else ""
        if uri.endswith(("/picture", "/chart", "/diagram")):
            drawings.append(drawing)
    return drawings


def _process_docx(file_bytes: bytes, filename: str) -> ExtractedDocument:
    document = open_docx(io.BytesIO(file_bytes))

    blocks: list[_Block] = []
    headings: list[HeadingItem] = []
    list_items = 0
    # (índice del bloque donde aparece, drawing, título de la sección vigente)
    drawings: list[tuple[int, object, str | None]] = []
    current_heading: str | None = None

    for item in document.iter_inner_content():
        if isinstance(item, Table):
            rows = []
            for row in item.rows:
                cells, seen = [], set()
                for cell in row.cells:
                    # Las celdas combinadas se repiten en row.cells: se toman una sola vez.
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    if cell.text.strip():
                        cells.append(" ".join(cell.text.split()))
                if cells:
                    rows.append(" | ".join(cells))
            index = len(blocks)
            if rows:
                blocks.append(_Block("\n".join(rows)))
            drawings.extend((index, d, current_heading) for d in _docx_drawings(item._tbl))
            continue

        paragraph: Paragraph = item
        text = paragraph.text.strip()
        level = _docx_heading_level(paragraph) if text else None
        index = len(blocks)
        if level:
            headings.append(
                HeadingItem(
                    level=level,
                    text=" ".join(text.split()),
                    page_number=1,
                    font_size=_docx_font_size(paragraph),
                    is_bold=_docx_is_bold(paragraph),
                )
            )
            current_heading = headings[-1].text
            blocks.append(_Block(text))
        elif text and _docx_is_list(paragraph):
            list_items += 1
            # Word dibuja la viñeta o el número, pero no forma parte del texto del párrafo.
            blocks.append(_Block(f"• {text}", is_list=True))
        elif text:
            blocks.append(_Block(text))
        drawings.extend((index, d, current_heading) for d in _docx_drawings(paragraph._p))

    images = _Images(total=len(drawings))
    for number, (index, drawing, section) in enumerate(drawings, start=1):
        has_alt, decorative = _alt_text_state(next(drawing.iter(w_qn("wp:docPr")), None))
        if has_alt or decorative:
            images.with_alt_text += 1
            continue

        extent = next(drawing.iter(w_qn("wp:extent")), None)
        width_pt = int(extent.get("cx", 0)) / _EMU_PER_PT if extent is not None else 0
        height_pt = int(extent.get("cy", 0)) / _EMU_PER_PT if extent is not None else 0
        if not _is_relevant_size(width_pt, height_pt):
            continue

        blob = None
        blip = next(drawing.iter(w_qn("a:blip")), None)
        r_id = blip.get(w_qn("r:embed")) if blip is not None else None
        if r_id and r_id in document.part.related_parts:
            blob = document.part.related_parts[r_id].blob

        # blocks[index] es el párrafo que contiene la imagen (si tiene texto) o el siguiente con texto.
        preceding = blocks[index - 1].text if index > 0 else None
        following = blocks[index].text if index < len(blocks) else None
        location = f"Imagen {number}, sección «{section}»" if section else f"Imagen {number} del documento"
        _add_candidate(
            images, blob, 1, (0.0, 0.0, width_pt, height_pt), location,
            _context_snippet(preceding, following),
        )

    # Estilo de párrafo por defecto ("Normal"; su nombre cambia según el idioma de Word).
    default_style = document.styles.default(WD_STYLE_TYPE.PARAGRAPH)
    normal_size = default_style.font.size if default_style is not None else None

    extracted = _build_document(
        "docx", filename, blocks, _docx_page_count(file_bytes), images,
        _structure(headings, list_items, normal_size.pt if normal_size is not None else None),
    )
    extracted.legibility = extract_docx_legibility(document)
    return extracted


def _docx_page_count(file_bytes: bytes) -> int:
    """Páginas según docProps/app.xml (Word lo guarda al último guardado); 1 si no está."""
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
            match = _PAGES_RE.search(zf.read("docProps/app.xml"))
    except KeyError:
        return 1
    return int(match.group(1)) if match else 1


# ---------------------------------------------------------------------------
# PowerPoint (.pptx)
# ---------------------------------------------------------------------------

_TITLE_PLACEHOLDERS = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
_BODY_PLACEHOLDERS = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.VERTICAL_BODY, PP_PLACEHOLDER.VERTICAL_OBJECT}


def _iter_shapes(shapes):
    """Formas en el orden del árbol de la diapositiva (el que usan los lectores de pantalla)."""
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _placeholder_type(shape):
    if not shape.is_placeholder:
        return None
    try:
        return shape.placeholder_format.type
    except ValueError:
        return None


def _pptx_is_list(paragraph, in_body_placeholder: bool) -> bool:
    """Viñeta explícita, o párrafo de un marcador de contenido (que hereda viñetas) sin buNone."""
    p_pr = paragraph._p.pPr
    if p_pr is not None:
        if p_pr.find(p_qn("a:buNone")) is not None:
            return False
        if p_pr.find(p_qn("a:buChar")) is not None or p_pr.find(p_qn("a:buAutoNum")) is not None:
            return True
    return in_body_placeholder or paragraph.level > 0


def _vertical_label(center: float, height: float) -> str:
    norm_y = center / max(height, 1.0)
    if norm_y < 0.35:
        return "sector superior"
    if norm_y > 0.65:
        return "sector inferior"
    return "sector central"


def _shape_box_pt(shape) -> tuple[float, float, float, float]:
    left, top = Emu(shape.left or 0).pt, Emu(shape.top or 0).pt
    return left, top, left + Emu(shape.width or 0).pt, top + Emu(shape.height or 0).pt


def _process_pptx(file_bytes: bytes, filename: str) -> ExtractedDocument:
    presentation = open_pptx(io.BytesIO(file_bytes))
    slide_height_pt = Emu(presentation.slide_height or 0).pt

    # Las diapositivas ocultas no se proyectan ni se comparten con los estudiantes.
    slides = [
        (number, slide)
        for number, slide in enumerate(presentation.slides, start=1)
        if slide._element.get("show") not in ("0", "false")
    ]
    if not slides:
        raise ValueError("La presentación de PowerPoint no contiene diapositivas visibles.")

    has_title_slide = any(
        _placeholder_type(shape) == PP_PLACEHOLDER.CENTER_TITLE
        for _, slide in slides
        for shape in slide.shapes
    )

    blocks: list[_Block] = []
    headings: list[HeadingItem] = []
    list_items = 0
    images = _Images()

    for number, slide in slides:
        slide_blocks: list[_Block] = []
        pictures = []

        for shape in _iter_shapes(slide.shapes):
            ph_type = _placeholder_type(shape)
            _, top, _, bottom = _shape_box_pt(shape)

            if isinstance(shape, Picture) or getattr(shape, "has_chart", False) or _is_smartart(shape):
                pictures.append(shape)
                continue

            if getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    cells = [" ".join(c.text.split()) for c in row.cells if c.text.strip()]
                    if cells:
                        rows.append(" | ".join(cells))
                if rows:
                    slide_blocks.append(_Block("\n".join(rows), top=top, bottom=bottom))
                continue

            if not shape.has_text_frame:
                continue

            if ph_type in _TITLE_PLACEHOLDERS:
                text = " ".join(shape.text_frame.text.split())
                if text:
                    level = 1 if ph_type == PP_PLACEHOLDER.CENTER_TITLE or not has_title_slide else 2
                    size = next(
                        (r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size),
                        None,
                    )
                    headings.append(HeadingItem(level=level, text=text, page_number=number, font_size=size))
                    slide_blocks.append(_Block(text, top=top, bottom=bottom))
                continue

            in_body = ph_type in _BODY_PLACEHOLDERS
            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.replace("\v", "\n").strip()
                if not text:
                    continue
                if _pptx_is_list(paragraph, in_body):
                    list_items += 1
                    slide_blocks.append(_Block(f"• {text}", is_list=True, top=top, bottom=bottom))
                else:
                    slide_blocks.append(_Block(text, top=top, bottom=bottom))

        images.total += len(pictures)
        for shape in pictures:
            has_alt, decorative = _alt_text_state(next(shape._element.iter(p_qn("p:cNvPr")), None))
            if has_alt or decorative:
                images.with_alt_text += 1
                continue

            x0, y0, x1, y1 = _shape_box_pt(shape)
            if not _is_relevant_size(x1 - x0, y1 - y0):
                continue

            blob = None
            if isinstance(shape, Picture):
                try:
                    blob = shape.image.blob
                except Exception:  # imagen vinculada a un archivo externo
                    blob = None

            preceding = [b.text for b in slide_blocks if b.bottom <= y0 + 5]
            following = [b.text for b in slide_blocks if b.top >= y1 - 5]
            location = f"Diapositiva {number}, {_vertical_label((y0 + y1) / 2, slide_height_pt)}"
            _add_candidate(
                images, blob, number, (x0, y0, x1, y1), location,
                _context_snippet(preceding[-1] if preceding else None, following[0] if following else None),
            )

        blocks.extend(slide_blocks)

    extracted = _build_document(
        "pptx", filename, blocks, len(slides), images, _structure(headings, list_items, None),
    )
    extracted.legibility = extract_pptx_legibility(presentation)
    return extracted


def _is_smartart(shape) -> bool:
    graphic_data = next(shape._element.iter(p_qn("a:graphicData")), None)
    return graphic_data is not None and graphic_data.get("uri", "").endswith("/diagram")


class OfficeDocumentService:
    """Procesa .docx y .pptx en memoria y los devuelve con el mismo contrato que un PDF."""

    @staticmethod
    def process_docx(file_bytes: bytes, filename: str) -> ExtractedDocument:
        return _wrap_errors("Word", _process_docx, file_bytes, filename)

    @staticmethod
    def process_pptx(file_bytes: bytes, filename: str) -> ExtractedDocument:
        return _wrap_errors("PowerPoint", _process_pptx, file_bytes, filename)


office_service = OfficeDocumentService()
