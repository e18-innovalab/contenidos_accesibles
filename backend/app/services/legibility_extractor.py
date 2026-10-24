"""
Extracción de datos tipográficos para las reglas de legibilidad visual (issue #79).

Hace su propia pasada sobre el archivo ya abierto para no modificar la extracción de texto
ni la detección de títulos de document_service.py y office_service.py. Todo es local y en
memoria: no consume tokens.

Por cada formato se mide:
- body_font_size: el tamaño de letra predominante, ponderado por caracteres.
- small_text_passages: párrafos (o bloques del PDF) con muchas palabras en letra chica.

Los fragmentos sin tamaño conocido (por ejemplo, texto de PowerPoint que hereda el tamaño
del patrón de diapositivas) no se cuentan: es preferible no marcar nada que inventar un tamaño.
"""

from docx.oxml.ns import qn as w_qn
from docx.table import Table
from pptx.oxml.ns import qn as p_qn
from pptx.shapes.group import GroupShape

from app.schemas.document import LegibilityStats, SmallTextPassage

# Un pasaje cuenta como "letra chica" si tiene al menos MIN_SMALL_PASSAGE_WORDS palabras en menos
# de SMALL_PASSAGE_PT. Así no se marcan números de página, pies de foto de dos palabras ni marcas
# de agua cortas.
SMALL_PASSAGE_PT = 8.0
MIN_SMALL_PASSAGE_WORDS = 20
MAX_STORED_PASSAGES = 20  # el total se informa aparte en small_text_passages_total
PASSAGE_EXCERPT_WORDS = 40


class _Collector:
    """Acumula tamaños por caracteres y arma la lista de pasajes con letra chica."""

    def __init__(self) -> None:
        self.char_counts: dict[float, int] = {}
        self.passages: list[SmallTextPassage] = []
        self.passages_total = 0

    def add_paragraph(self, fragments: list[tuple[str, float | None]], page_number: int) -> None:
        """fragments: (texto, tamaño en pt o None si no se conoce) de un mismo párrafo o bloque."""
        small_texts: list[str] = []
        small_sizes: dict[float, int] = {}
        small_words = 0
        for text, size in fragments:
            stripped = text.strip()
            if not stripped or size is None:
                continue
            size = round(size, 1)
            self.char_counts[size] = self.char_counts.get(size, 0) + len(stripped)
            if size < SMALL_PASSAGE_PT:
                small_texts.append(stripped)
                small_sizes[size] = small_sizes.get(size, 0) + len(stripped)
                small_words += len(stripped.split())

        if small_words < MIN_SMALL_PASSAGE_WORDS:
            return
        self.passages_total += 1
        if len(self.passages) >= MAX_STORED_PASSAGES:
            return
        words = " ".join(small_texts).split()
        excerpt = " ".join(words[:PASSAGE_EXCERPT_WORDS]) + ("…" if len(words) > PASSAGE_EXCERPT_WORDS else "")
        self.passages.append(
            SmallTextPassage(
                text=excerpt,
                page_number=max(page_number, 1),
                font_size=max(small_sizes, key=small_sizes.get),
                word_count=small_words,
            )
        )

    def result(self) -> LegibilityStats:
        return LegibilityStats(
            body_font_size=max(self.char_counts, key=self.char_counts.get) if self.char_counts else None,
            measured_chars=sum(self.char_counts.values()),
            small_text_passages=self.passages,
            small_text_passages_total=self.passages_total,
        )


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def extract_pdf_legibility(doc) -> LegibilityStats:
    """doc es un pymupdf.Document ya abierto. Cada bloque de texto se trata como un párrafo."""
    collector = _Collector()
    for page_idx in range(len(doc)):
        page_dict = doc.load_page(page_idx).get_text("dict")
        for block in page_dict.get("blocks", []):
            if block.get("type") != 0:  # 0 = texto
                continue
            fragments = [
                (span.get("text", ""), span.get("size"))
                for line in block.get("lines", [])
                for span in line.get("spans", [])
            ]
            collector.add_paragraph(fragments, page_idx + 1)
    return collector.result()


# ---------------------------------------------------------------------------
# Word (.docx)
# ---------------------------------------------------------------------------

def _style_sizes(style):
    """Tamaño definido en el estilo o en alguno de sus estilos base."""
    while style is not None:
        if style.font.size is not None:
            return style.font.size.pt
        style = style.base_style
    return None


def _docx_default_size(document) -> float | None:
    """Tamaño de docDefaults (w:sz está en medios puntos). Word usa 10 pt si no está definido."""
    sz = document.styles.element.find(
        f"{w_qn('w:docDefaults')}/{w_qn('w:rPrDefault')}/{w_qn('w:rPr')}/{w_qn('w:sz')}"
    )
    if sz is not None and sz.get(w_qn("w:val"), "").isdigit():
        return int(sz.get(w_qn("w:val"))) / 2
    return 10.0


def _docx_run_size(run, paragraph, default_size: float | None) -> float | None:
    if run.font.size is not None:
        return run.font.size.pt
    for style in (run.style, paragraph.style):
        size = _style_sizes(style) if style is not None else None
        if size is not None:
            return size
    return default_size


def _docx_paragraphs(document):
    """Párrafos del cuerpo y de las celdas de tabla (las celdas combinadas, una sola vez)."""
    for item in document.iter_inner_content():
        if isinstance(item, Table):
            seen = set()
            for row in item.rows:
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    yield from cell.paragraphs
        else:
            yield item


def extract_docx_legibility(document) -> LegibilityStats:
    """document es un docx.Document ya abierto. Word no informa páginas: todo va como página 1."""
    collector = _Collector()
    default_size = _docx_default_size(document)
    for paragraph in _docx_paragraphs(document):
        fragments = [(run.text, _docx_run_size(run, paragraph, default_size)) for run in paragraph.runs]
        collector.add_paragraph(fragments, 1)
    return collector.result()


# ---------------------------------------------------------------------------
# PowerPoint (.pptx)
# ---------------------------------------------------------------------------

def _pptx_font_scale(text_frame) -> float:
    """Factor de 'Reducir el texto al desbordar' (normAutofit fontScale, en milésimas de %)."""
    autofit = text_frame._bodyPr.find(p_qn("a:normAutofit"))
    scale = autofit.get("fontScale") if autofit is not None else None
    return int(scale) / 100_000 if scale and scale.isdigit() else 1.0


def _pptx_text_frames(shapes):
    """Marcos de texto de las formas, de las formas agrupadas y de las celdas de tabla."""
    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from _pptx_text_frames(shape.shapes)
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame
        elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            yield shape.text_frame


def extract_pptx_legibility(presentation) -> LegibilityStats:
    """presentation es un pptx.Presentation ya abierto. Se saltean las diapositivas ocultas."""
    collector = _Collector()
    for number, slide in enumerate(presentation.slides, start=1):
        if slide._element.get("show") in ("0", "false"):
            continue
        for text_frame in _pptx_text_frames(slide.shapes):
            scale = _pptx_font_scale(text_frame)
            for paragraph in text_frame.paragraphs:
                fragments = []
                for run in paragraph.runs:
                    size = run.font.size or paragraph.font.size
                    fragments.append((run.text, size.pt * scale if size is not None else None))
                collector.add_paragraph(fragments, number)
    return collector.result()
